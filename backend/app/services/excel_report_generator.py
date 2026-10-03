import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.services.data_ingestion import load_dataframe
from app.services.kpi_engine import compute_dataset_kpis
from app.services.insight_engine import generate_ai_business_insights

def generate_excel_report(
    file_path: str,
    output_excel_path: str,
    transformation_log: list = None
) -> str:
    
    os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)
    df = load_dataframe(file_path)
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    blue_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    bold_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=11, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # --- TAB 1: EXECUTIVE SUMMARY ---
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum.cell(row=2, column=2, value="AUTOMATED BI EXECUTIVE REPORT").font = title_font
    
    # Write KPIs
    ws_sum.cell(row=4, column=2, value="Key Performance Indicators").font = Font(name="Calibri", size=13, bold=True, color="2563EB")
    headers_kpi = ["Metric Name", "Value", "Trend", "Status"]
    for c_idx, h in enumerate(headers_kpi, start=2):
        cell = ws_sum.cell(row=5, column=c_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    kpi_res = compute_dataset_kpis(file_path)
    kpis = kpi_res.get("kpis", [])
    row_idx = 6
    for k in kpis:
        ws_sum.cell(row=row_idx, column=2, value=k.get("name")).font = bold_font
        ws_sum.cell(row=row_idx, column=3, value=k.get("value")).font = regular_font
        ws_sum.cell(row=row_idx, column=4, value=k.get("trend")).font = regular_font
        ws_sum.cell(row=row_idx, column=5, value=k.get("status").upper()).font = bold_font
        for c in range(2, 6):
            ws_sum.cell(row=row_idx, column=c).border = thin_border
            if row_idx % 2 == 1:
                ws_sum.cell(row=row_idx, column=c).fill = gray_fill
        row_idx += 1

    # Write AI Insights
    row_idx += 2
    ws_sum.cell(row=row_idx, column=2, value="AI Strategic Insights").font = Font(name="Calibri", size=13, bold=True, color="2563EB")
    row_idx += 1
    headers_ins = ["Category", "Insight Title", "Executive Description", "Recommended Action"]
    for c_idx, h in enumerate(headers_ins, start=2):
        cell = ws_sum.cell(row=row_idx, column=c_idx, value=h)
        cell.fill = blue_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    row_idx += 1
    
    insights = generate_ai_business_insights(file_path)
    for ins in insights:
        ws_sum.cell(row=row_idx, column=2, value=ins.get("category")).font = bold_font
        ws_sum.cell(row=row_idx, column=3, value=ins.get("title")).font = bold_font
        ws_sum.cell(row=row_idx, column=4, value=ins.get("description")).font = regular_font
        ws_sum.cell(row=row_idx, column=5, value=ins.get("recommendation")).font = regular_font
        for c in range(2, 6):
            ws_sum.cell(row=row_idx, column=c).border = thin_border
            if row_idx % 2 == 1:
                ws_sum.cell(row=row_idx, column=c).fill = gray_fill
        row_idx += 1

    # Adjust widths for Executive Summary
    for col in ws_sum.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_sum.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # --- TAB 2: RAW DATASET ---
    ws_data = wb.create_sheet(title="Raw Data")
    ws_data.views.sheetView[0].showGridLines = True
    
    # Write Headers
    for c_idx, col_name in enumerate(df.columns, start=1):
        cell = ws_data.cell(row=1, column=c_idx, value=str(col_name))
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_data.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if r_idx % 2 == 1:
                cell.fill = gray_fill
                
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

    # --- TAB 3: DATA QUALITY AUDIT ---
    if transformation_log:
        ws_audit = wb.create_sheet(title="Transformation Audit")
        ws_audit.views.sheetView[0].showGridLines = True
        
        ws_audit.cell(row=1, column=1, value="Step #").fill = navy_fill
        ws_audit.cell(row=1, column=2, value="Column / Target").fill = navy_fill
        ws_audit.cell(row=1, column=3, value="Action Applied").fill = navy_fill
        ws_audit.cell(row=1, column=4, value="Detailed Log").fill = navy_fill
        for c in range(1, 5):
            ws_audit.cell(row=1, column=c).font = header_font
            
        for r_idx, t in enumerate(transformation_log, start=2):
            ws_audit.cell(row=r_idx, column=1, value=t.get("step")).font = bold_font
            ws_audit.cell(row=r_idx, column=2, value=t.get("column", "Dataset Level")).font = regular_font
            ws_audit.cell(row=r_idx, column=3, value=t.get("action")).font = bold_font
            ws_audit.cell(row=r_idx, column=4, value=t.get("detail")).font = regular_font
            for c in range(1, 5):
                ws_audit.cell(row=r_idx, column=c).border = thin_border
                
        for col in ws_audit.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_audit.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    wb.save(output_excel_path)
    return output_excel_path
